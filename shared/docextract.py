"""添付ドキュメントのテキスト抽出（backend / rag-app 共通モジュール）。

ローカル LLM は PDF/Word/Excel 等を直接読めないため、テキストへ抽出して
プロンプトや RAG の知識ベースに流し込むための共通関数を提供する。

依存: pypdf / python-docx / openpyxl（各サービスの requirements に含める）。
"""

from __future__ import annotations

import base64
import io
import os

# 抽出テキストの最大長（コンテキスト肥大を防ぐ）
MAX_DOC_CHARS = int(os.environ.get("MAX_DOC_CHARS", "30000"))

# UI の accept などに使える対応拡張子
SUPPORTED_DOC_EXTS = (
    ".pdf",
    ".docx",
    ".xlsx",
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".html",
    ".htm",
    ".json",
    ".log",
)


def strip_base64_prefix(data: str) -> str:
    """`data:application/pdf;base64,xxxx` のような prefix を除去する。"""
    if data.startswith("data:"):
        comma = data.find(",")
        if comma != -1:
            return data[comma + 1 :]
    return data


def b64_to_bytes(data: str) -> bytes:
    return base64.b64decode(strip_base64_prefix(data))


def extract_doc_text(name: str, media_type: str, b64: str) -> str | None:
    """添付ドキュメント(PDF/Word/Excel/テキスト)からテキストを抽出する。

    対応外（レガシー .doc/.xls 等）は None を返す。
    """
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    mt = media_type or ""
    try:
        raw = b64_to_bytes(b64)
    except Exception:  # noqa: BLE001
        return None

    try:
        if ext == "pdf" or mt == "application/pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            text = "\n".join((p.extract_text() or "") for p in reader.pages)
        elif ext == "docx" or "wordprocessingml" in mt:
            import docx

            d = docx.Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in d.paragraphs)
        elif ext == "xlsx" or "spreadsheetml" in mt:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join("" if c is None else str(c) for c in row))
            text = "\n".join(lines)
        elif ext in (
            "txt",
            "md",
            "csv",
            "tsv",
            "html",
            "htm",
            "json",
            "log",
        ) or mt.startswith("text/"):
            text = raw.decode("utf-8", "ignore")
        else:
            return None
    except Exception as e:  # noqa: BLE001
        return f"(添付ファイル {name} のテキスト抽出に失敗しました: {e})"

    text = (text or "").strip()
    if not text:
        return f"(添付ファイル {name} からテキストを抽出できませんでした)"
    if len(text) > MAX_DOC_CHARS:
        text = text[:MAX_DOC_CHARS] + "\n…(以下省略)"
    return text
